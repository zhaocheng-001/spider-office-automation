# -*- coding: utf-8 -*-
"""
openpyxl Excel 报表工具：
- write_report：把二维数据写成一个带样式的 Excel 文件（表头加粗填色、自适应列宽、¥ 金额格式）
- read_excel：读取 Excel 为二维列表
- generate_all_reports：从 MySQL 取数，批量生成多张报表到 reports/ 目录
"""
import logging
import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import Config
from db import get_connection

logger = logging.getLogger("office.excel")

HEADER_FILL = PatternFill("solid", fgColor="185FA5")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def write_report(headers, rows, filepath, title="报表", money_cols=()):
    """
    生成单个 Excel 报表。
    :param headers: 表头列表，如 ["书名", "价格"]
    :param rows: 数据行（二维列表）
    :param filepath: 输出 xlsx 路径
    :param money_cols: 需要 ¥ 千分位格式的列下标（从 0 开始）
    """
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]  # sheet 名最长 31 字符

    # 表头
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    for row in rows:
        ws.append(list(row))

    # 金额列格式：¥ + 千分位
    for col_idx in money_cols:
        col = get_column_letter(col_idx + 1)
        for c in ws[col][1:]:
            c.number_format = '"¥"#,##0.00'

    # 自适应列宽（按每列最长文本估算）
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len * 1.8 + 4, 10), 60)

    ws.freeze_panes = "A2"  # 冻结表头
    wb.save(filepath)
    logger.info("报表已生成：%s", filepath)
    return filepath


def read_excel(filepath, sheet=0, has_header=True):
    """读取 Excel：返回 (表头列表, 数据二维列表)"""
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    if has_header and rows:
        return rows[0], rows[1:]
    return None, rows


def _fetch_quotes_stats():
    """从 MySQL 统计每个作者的语录数量"""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT author, COUNT(*) AS cnt, MAX(crawled_at) "
                "FROM quotes GROUP BY author ORDER BY cnt DESC"
            )
            return cur.fetchall()
    finally:
        conn.close()


def _fetch_books():
    """从 MySQL 取图书列表"""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT title, price, rating, in_stock FROM books ORDER BY price DESC")
            return cur.fetchall()
    finally:
        conn.close()


def generate_all_reports():
    """
    批量生成报表：reports/ 目录下按日期生成
    1. 语录统计报表（按作者分组）
    2. 图书价格报表（¥ 金额格式演示）
    返回生成的文件路径列表。
    """
    out_dir = os.path.join(Config.REPORT_DIR, datetime.now().strftime("%Y%m%d"))
    paths = []

    quotes_stats = _fetch_quotes_stats()
    if quotes_stats:
        paths.append(write_report(
            headers=["作者", "语录数量", "最近抓取时间"],
            rows=quotes_stats,
            filepath=os.path.join(out_dir, "语录作者统计报表.xlsx"),
            title="语录统计",
        ))

    books = _fetch_books()
    if books:
        paths.append(write_report(
            headers=["书名", "价格", "评分", "有货"],
            rows=books,
            filepath=os.path.join(out_dir, "图书价格报表.xlsx"),
            title="图书价格",
            money_cols=(1,),  # 价格列加 ¥ 格式
        ))

    if not paths:
        logger.warning("数据库没有数据，先运行爬虫：python -m spiders.quote_spider")
    return paths
